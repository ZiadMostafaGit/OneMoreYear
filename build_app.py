#!/usr/bin/env python3
"""Read the Excel schedule and generate a self-contained HTML study dashboard."""

import json
import datetime
from openpyxl import load_workbook

def extract_data(xlsx_path):
    wb = load_workbook(xlsx_path)
    ws = wb["365-Day Study Plan"]

    days = {}
    current_day = None

    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, values_only=True):
        day_num, date_str, phase, total_hours, domain, topic, subtopic, duration, difficulty, notes = row

        if day_num and day_num != "":
            current_day = int(day_num)
            days[current_day] = {
                "day": current_day,
                "date": str(date_str) if date_str else "",
                "phase": phase if phase else "",
                "totalHours": float(total_hours) if total_hours else 0,
                "blocks": []
            }

        if current_day and current_day in days:
            diff_val = 0
            if difficulty:
                diff_val = len(str(difficulty))
            days[current_day]["blocks"].append({
                "domain": domain or "",
                "topic": topic or "",
                "subtopic": subtopic or "",
                "duration": float(duration) if duration else 0,
                "difficulty": diff_val,
                "notes": notes or ""
            })

    return list(days.values())


def build_html(data):
    data_json = json.dumps(data, ensure_ascii=False)

    return f'''<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0">
<title>FAANG Backend Engineering — 365-Day Study Plan</title>
<style>
* {{ margin: 0; padding: 0; box-sizing: border-box; }}

:root {{
  --bg: #0f1117;
  --surface: #1a1d27;
  --surface2: #232733;
  --border: #2e3345;
  --text: #e4e6ed;
  --text2: #9098b0;
  --accent: #6c8cff;
  --accent2: #4a6cf7;
  --green: #34d399;
  --green-dim: #065f46;
  --yellow: #fbbf24;
  --orange: #f97316;
  --red: #ef4444;
  --purple: #a78bfa;
  --pink: #f472b6;
  --cyan: #22d3ee;
  --foundation: #34d399;
  --intermediate: #6c8cff;
  --advanced: #f97316;
  --integration: #a78bfa;
}}

body {{
  font-family: -apple-system, BlinkMacSystemFont, 'Segoe UI', system-ui, sans-serif;
  background: var(--bg);
  color: var(--text);
  min-height: 100vh;
  line-height: 1.5;
}}

/* ── Top Bar ── */
.topbar {{
  background: var(--surface);
  border-bottom: 1px solid var(--border);
  padding: 14px 24px;
  display: flex;
  align-items: center;
  justify-content: space-between;
  position: sticky;
  top: 0;
  z-index: 100;
  backdrop-filter: blur(12px);
}}
.topbar h1 {{
  font-size: 18px;
  font-weight: 700;
  background: linear-gradient(135deg, var(--accent), var(--purple));
  -webkit-background-clip: text;
  -webkit-text-fill-color: transparent;
}}
.topbar-stats {{
  display: flex;
  gap: 20px;
  font-size: 13px;
  color: var(--text2);
}}
.topbar-stats strong {{ color: var(--text); }}

/* ── Layout ── */
.container {{
  max-width: 1100px;
  margin: 0 auto;
  padding: 24px;
}}

/* ── Phase Progress ── */
.progress-section {{
  display: grid;
  grid-template-columns: repeat(4, 1fr);
  gap: 12px;
  margin-bottom: 28px;
}}
.phase-card {{
  background: var(--surface);
  border-radius: 12px;
  padding: 16px;
  border: 1px solid var(--border);
  cursor: pointer;
  transition: all 0.2s;
}}
.phase-card:hover {{ border-color: var(--accent); transform: translateY(-1px); }}
.phase-card.active {{ border-color: var(--accent); box-shadow: 0 0 20px rgba(108,140,255,0.15); }}
.phase-label {{
  font-size: 11px;
  text-transform: uppercase;
  letter-spacing: 1.2px;
  font-weight: 600;
  margin-bottom: 6px;
}}
.phase-label.foundation {{ color: var(--foundation); }}
.phase-label.intermediate {{ color: var(--intermediate); }}
.phase-label.advanced {{ color: var(--advanced); }}
.phase-label.integration {{ color: var(--integration); }}
.phase-days {{ font-size: 13px; color: var(--text2); margin-bottom: 8px; }}
.progress-bar {{
  height: 6px;
  background: var(--surface2);
  border-radius: 3px;
  overflow: hidden;
}}
.progress-fill {{
  height: 100%;
  border-radius: 3px;
  transition: width 0.5s ease;
}}
.progress-fill.foundation {{ background: var(--foundation); }}
.progress-fill.intermediate {{ background: var(--intermediate); }}
.progress-fill.advanced {{ background: var(--advanced); }}
.progress-fill.integration {{ background: var(--integration); }}
.phase-pct {{
  font-size: 22px;
  font-weight: 700;
  margin-top: 6px;
}}

/* ── Navigation ── */
.nav-bar {{
  display: flex;
  align-items: center;
  justify-content: space-between;
  margin-bottom: 20px;
  gap: 12px;
  flex-wrap: wrap;
}}
.nav-group {{
  display: flex;
  align-items: center;
  gap: 8px;
}}
.nav-btn {{
  background: var(--surface);
  border: 1px solid var(--border);
  color: var(--text);
  padding: 8px 16px;
  border-radius: 8px;
  cursor: pointer;
  font-size: 14px;
  transition: all 0.15s;
}}
.nav-btn:hover {{ background: var(--surface2); border-color: var(--accent); }}
.nav-btn.today {{ background: var(--accent2); border-color: var(--accent); font-weight: 600; }}
.day-title {{
  font-size: 28px;
  font-weight: 700;
}}
.day-meta {{
  font-size: 14px;
  color: var(--text2);
  display: flex;
  gap: 16px;
  align-items: center;
}}
.day-meta .phase-badge {{
  padding: 3px 10px;
  border-radius: 12px;
  font-size: 12px;
  font-weight: 600;
  text-transform: uppercase;
  letter-spacing: 0.5px;
}}
.phase-badge.foundation {{ background: rgba(52,211,153,0.15); color: var(--foundation); }}
.phase-badge.intermediate {{ background: rgba(108,140,255,0.15); color: var(--intermediate); }}
.phase-badge.advanced {{ background: rgba(249,115,22,0.15); color: var(--advanced); }}
.phase-badge.integration {{ background: rgba(167,139,250,0.15); color: var(--integration); }}

/* ── Day Card ── */
.day-progress {{
  background: var(--surface);
  border-radius: 12px;
  padding: 16px 20px;
  margin-bottom: 16px;
  border: 1px solid var(--border);
  display: flex;
  align-items: center;
  justify-content: space-between;
}}
.day-progress-bar {{
  flex: 1;
  margin: 0 20px;
  height: 8px;
  background: var(--surface2);
  border-radius: 4px;
  overflow: hidden;
}}
.day-progress-fill {{
  height: 100%;
  background: var(--green);
  border-radius: 4px;
  transition: width 0.3s;
}}
.day-pct {{ font-weight: 700; font-size: 18px; min-width: 50px; text-align: right; }}
.day-hours {{ color: var(--text2); font-size: 14px; }}

/* ── Study Blocks ── */
.blocks-list {{ display: flex; flex-direction: column; gap: 8px; }}

.study-block {{
  background: var(--surface);
  border: 1px solid var(--border);
  border-radius: 12px;
  padding: 16px 20px;
  display: flex;
  align-items: flex-start;
  gap: 16px;
  transition: all 0.2s;
  cursor: pointer;
}}
.study-block:hover {{ border-color: var(--border); background: var(--surface2); }}
.study-block.done {{ opacity: 0.5; }}
.study-block.done .block-subtopic {{ text-decoration: line-through; color: var(--text2); }}

.block-check {{
  flex-shrink: 0;
  width: 22px;
  height: 22px;
  border: 2px solid var(--border);
  border-radius: 6px;
  cursor: pointer;
  display: flex;
  align-items: center;
  justify-content: center;
  margin-top: 2px;
  transition: all 0.15s;
}}
.block-check:hover {{ border-color: var(--accent); }}
.block-check.checked {{
  background: var(--green);
  border-color: var(--green);
}}
.block-check.checked::after {{
  content: "\\2713";
  color: white;
  font-size: 14px;
  font-weight: 700;
}}

.block-body {{ flex: 1; min-width: 0; }}
.block-header {{ display: flex; align-items: center; gap: 10px; margin-bottom: 4px; flex-wrap: wrap; }}

.domain-tag {{
  font-size: 11px;
  padding: 2px 8px;
  border-radius: 6px;
  font-weight: 600;
  white-space: nowrap;
}}
.domain-tag.computer-fundamentals {{ background: rgba(34,211,238,0.15); color: var(--cyan); }}
.domain-tag.operating-systems {{ background: rgba(52,211,153,0.15); color: var(--foundation); }}
.domain-tag.networking {{ background: rgba(108,140,255,0.15); color: var(--intermediate); }}
.domain-tag.linux-operations {{ background: rgba(251,191,36,0.15); color: var(--yellow); }}
.domain-tag.apis-protocols {{ background: rgba(249,115,22,0.15); color: var(--orange); }}
.domain-tag.design-patterns {{ background: rgba(167,139,250,0.15); color: var(--purple); }}
.domain-tag.databases {{ background: rgba(244,114,182,0.15); color: var(--pink); }}
.domain-tag.devops {{ background: rgba(251,191,36,0.15); color: var(--yellow); }}
.domain-tag.distributed-systems {{ background: rgba(239,68,68,0.15); color: var(--red); }}
.domain-tag.infrastructure {{ background: rgba(34,211,238,0.15); color: var(--cyan); }}
.domain-tag.advanced-architecture {{ background: rgba(167,139,250,0.15); color: var(--purple); }}
.domain-tag.system-design {{ background: rgba(108,140,255,0.15); color: var(--accent); }}
.domain-tag.review {{ background: rgba(251,191,36,0.15); color: var(--yellow); }}

.block-topic {{ font-size: 12px; color: var(--text2); }}
.block-subtopic {{ font-size: 15px; font-weight: 500; line-height: 1.4; }}
.block-footer {{ display: flex; gap: 12px; margin-top: 6px; align-items: center; }}
.block-duration {{
  font-size: 12px;
  color: var(--text2);
  background: var(--surface2);
  padding: 2px 8px;
  border-radius: 4px;
}}
.block-difficulty {{ font-size: 11px; letter-spacing: 1px; }}
.block-notes {{ font-size: 12px; color: var(--text2); font-style: italic; margin-top: 4px; }}

/* ── Day picker overlay ── */
.day-picker-overlay {{
  display: none;
  position: fixed;
  inset: 0;
  background: rgba(0,0,0,0.6);
  z-index: 200;
  align-items: center;
  justify-content: center;
}}
.day-picker-overlay.open {{ display: flex; }}
.day-picker {{
  background: var(--surface);
  border: 1px solid var(--border);
  border-radius: 16px;
  padding: 24px;
  max-width: 720px;
  width: 95%;
  max-height: 80vh;
  overflow-y: auto;
}}
.day-picker h2 {{ margin-bottom: 16px; }}
.day-grid {{
  display: grid;
  grid-template-columns: repeat(auto-fill, minmax(38px, 1fr));
  gap: 4px;
}}
.day-cell {{
  width: 100%;
  aspect-ratio: 1;
  border-radius: 6px;
  border: 1px solid var(--border);
  display: flex;
  align-items: center;
  justify-content: center;
  font-size: 11px;
  cursor: pointer;
  transition: all 0.1s;
  color: var(--text2);
}}
.day-cell:hover {{ border-color: var(--accent); color: var(--text); }}
.day-cell.current {{ border-color: var(--accent); background: var(--accent2); color: white; font-weight: 700; }}
.day-cell.completed {{ background: var(--green-dim); border-color: var(--green); color: var(--green); }}
.day-cell.partial {{ background: rgba(251,191,36,0.1); border-color: var(--yellow); color: var(--yellow); }}

/* ── Jump input ── */
.jump-input {{
  background: var(--surface);
  border: 1px solid var(--border);
  color: var(--text);
  padding: 8px 12px;
  border-radius: 8px;
  font-size: 14px;
  width: 70px;
  text-align: center;
}}
.jump-input:focus {{ outline: none; border-color: var(--accent); }}

/* ── Responsive ── */
@media (max-width: 768px) {{
  .progress-section {{ grid-template-columns: repeat(2, 1fr); }}
  .topbar {{ flex-direction: column; gap: 8px; }}
  .nav-bar {{ flex-direction: column; align-items: flex-start; }}
  .container {{ padding: 16px; }}
  .day-title {{ font-size: 22px; }}
}}
@media (max-width: 480px) {{
  .progress-section {{ grid-template-columns: 1fr; }}
}}
</style>
</head>
<body>

<div class="topbar">
  <h1>FAANG Backend Engineering — 365-Day Plan</h1>
  <div class="topbar-stats">
    <span><strong id="stat-completed">0</strong> / 365 days</span>
    <span><strong id="stat-hours">0</strong>h studied</span>
    <span><strong id="stat-streak">0</strong> day streak</span>
  </div>
</div>

<div class="container">
  <!-- Phase Progress -->
  <div class="progress-section">
    <div class="phase-card" data-phase="Foundation" onclick="jumpToPhase('Foundation')">
      <div class="phase-label foundation">Phase 1 — Foundation</div>
      <div class="phase-days">Days 1–95</div>
      <div class="progress-bar"><div class="progress-fill foundation" id="prog-foundation" style="width:0%"></div></div>
      <div class="phase-pct" id="pct-foundation">0%</div>
    </div>
    <div class="phase-card" data-phase="Intermediate" onclick="jumpToPhase('Intermediate')">
      <div class="phase-label intermediate">Phase 2 — Intermediate</div>
      <div class="phase-days">Days 96–195</div>
      <div class="progress-bar"><div class="progress-fill intermediate" id="prog-intermediate" style="width:0%"></div></div>
      <div class="phase-pct" id="pct-intermediate">0%</div>
    </div>
    <div class="phase-card" data-phase="Advanced" onclick="jumpToPhase('Advanced')">
      <div class="phase-label advanced">Phase 3 — Advanced</div>
      <div class="phase-days">Days 196–290</div>
      <div class="progress-bar"><div class="progress-fill advanced" id="prog-advanced" style="width:0%"></div></div>
      <div class="phase-pct" id="pct-advanced">0%</div>
    </div>
    <div class="phase-card" data-phase="Integration" onclick="jumpToPhase('Integration')">
      <div class="phase-label integration">Phase 4 — Integration</div>
      <div class="phase-days">Days 291–365</div>
      <div class="progress-bar"><div class="progress-fill integration" id="prog-integration" style="width:0%"></div></div>
      <div class="phase-pct" id="pct-integration">0%</div>
    </div>
  </div>

  <!-- Navigation -->
  <div class="nav-bar">
    <div>
      <div class="day-title" id="day-title">Day 1</div>
      <div class="day-meta">
        <span id="day-date"></span>
        <span class="phase-badge" id="day-phase-badge"></span>
      </div>
    </div>
    <div class="nav-group">
      <button class="nav-btn" onclick="prevDay()">&#9664; Prev</button>
      <button class="nav-btn today" onclick="goToday()">Today</button>
      <button class="nav-btn" onclick="nextDay()">Next &#9654;</button>
      <input type="number" class="jump-input" id="jump-input" min="1" max="365" placeholder="Day" onchange="jumpTo(this.value)">
      <button class="nav-btn" onclick="openPicker()">&#9783; Calendar</button>
    </div>
  </div>

  <!-- Day Progress -->
  <div class="day-progress">
    <div class="day-hours"><span id="done-hours">0</span> / <span id="total-hours">0</span>h</div>
    <div class="day-progress-bar"><div class="day-progress-fill" id="day-prog-fill"></div></div>
    <div class="day-pct" id="day-pct">0%</div>
  </div>

  <!-- Blocks -->
  <div class="blocks-list" id="blocks-list"></div>
</div>

<!-- Day Picker -->
<div class="day-picker-overlay" id="picker-overlay" onclick="if(event.target===this)closePicker()">
  <div class="day-picker">
    <h2>Jump to Day</h2>
    <div class="day-grid" id="day-grid"></div>
  </div>
</div>

<script>
const DATA = {data_json};

// State
const START_DATE = new Date("2026-04-13");
let currentDay = 1;
let checked = JSON.parse(localStorage.getItem("faang365_checked") || "{{}}");

// Helpers
function domainClass(d) {{
  return d.toLowerCase().replace(/[^a-z]/g, '-').replace(/-+/g, '-').replace(/-$/, '');
}}

function getDayData(n) {{
  return DATA.find(d => d.day === n);
}}

function blockKey(dayNum, idx) {{
  return dayNum + "_" + idx;
}}

function isBlockChecked(dayNum, idx) {{
  return !!checked[blockKey(dayNum, idx)];
}}

function toggleBlock(dayNum, idx) {{
  const k = blockKey(dayNum, idx);
  if (checked[k]) delete checked[k]; else checked[k] = true;
  localStorage.setItem("faang365_checked", JSON.stringify(checked));
  renderDay(currentDay);
  updateStats();
}}

function dayCompletion(dayNum) {{
  const d = getDayData(dayNum);
  if (!d) return 0;
  const total = d.blocks.length;
  if (total === 0) return 0;
  let done = 0;
  d.blocks.forEach((_, i) => {{ if (isBlockChecked(dayNum, i)) done++; }});
  return done / total;
}}

function calcToday() {{
  const now = new Date();
  const diff = Math.floor((now - START_DATE) / 86400000) + 1;
  return Math.max(1, Math.min(365, diff));
}}

// Render
function renderDay(n) {{
  const d = getDayData(n);
  if (!d) return;

  document.getElementById("day-title").textContent = "Day " + d.day;
  document.getElementById("day-date").textContent = d.date;
  const badge = document.getElementById("day-phase-badge");
  badge.textContent = d.phase;
  badge.className = "phase-badge " + d.phase.toLowerCase();

  document.getElementById("total-hours").textContent = d.totalHours.toFixed(1);
  document.getElementById("jump-input").value = n;

  // Blocks
  const list = document.getElementById("blocks-list");
  list.innerHTML = "";
  let doneHrs = 0;
  d.blocks.forEach((b, i) => {{
    const isDone = isBlockChecked(n, i);
    if (isDone) doneHrs += b.duration;

    const el = document.createElement("div");
    el.className = "study-block" + (isDone ? " done" : "");

    const dc = domainClass(b.domain);
    const stars = "\\u2605".repeat(b.difficulty);

    el.innerHTML = `
      <div class="block-check ${{isDone ? 'checked' : ''}}" onclick="event.stopPropagation();toggleBlock(${{n}},${{i}})"></div>
      <div class="block-body">
        <div class="block-header">
          <span class="domain-tag ${{dc}}">${{b.domain}}</span>
          <span class="block-topic">${{b.topic}}</span>
        </div>
        <div class="block-subtopic">${{b.subtopic}}</div>
        <div class="block-footer">
          <span class="block-duration">${{b.duration}}h</span>
          <span class="block-difficulty">${{stars}}</span>
        </div>
        ${{b.notes ? '<div class="block-notes">' + b.notes + '</div>' : ''}}
      </div>
    `;
    el.addEventListener("click", () => toggleBlock(n, i));
    list.appendChild(el);
  }});

  // Day progress
  const pct = d.blocks.length ? Math.round(doneHrs / d.totalHours * 100) : 0;
  document.getElementById("done-hours").textContent = doneHrs.toFixed(1);
  document.getElementById("day-prog-fill").style.width = Math.min(100, pct) + "%";
  document.getElementById("day-pct").textContent = pct + "%";
}}

function updateStats() {{
  let completedDays = 0;
  let totalHours = 0;
  let streak = 0;
  const today = calcToday();

  // Phase stats
  const phases = {{
    "Foundation": {{ start: 1, end: 95, done: 0, total: 0 }},
    "Intermediate": {{ start: 96, end: 195, done: 0, total: 0 }},
    "Advanced": {{ start: 196, end: 290, done: 0, total: 0 }},
    "Integration": {{ start: 291, end: 365, done: 0, total: 0 }},
  }};

  DATA.forEach(d => {{
    const comp = dayCompletion(d.day);
    if (comp >= 1) completedDays++;

    d.blocks.forEach((b, i) => {{
      if (isBlockChecked(d.day, i)) totalHours += b.duration;
    }});

    // Phase tracking
    for (const [name, p] of Object.entries(phases)) {{
      if (d.day >= p.start && d.day <= p.end) {{
        p.total++;
        if (comp >= 1) p.done++;
      }}
    }}
  }});

  // Streak
  for (let day = today; day >= 1; day--) {{
    if (dayCompletion(day) >= 1) streak++;
    else break;
  }}

  document.getElementById("stat-completed").textContent = completedDays;
  document.getElementById("stat-hours").textContent = Math.round(totalHours);
  document.getElementById("stat-streak").textContent = streak;

  // Phase bars
  for (const [name, p] of Object.entries(phases)) {{
    const pct = p.total ? Math.round(p.done / p.total * 100) : 0;
    const key = name.toLowerCase();
    document.getElementById("prog-" + key).style.width = pct + "%";
    document.getElementById("pct-" + key).textContent = pct + "%";
  }}
}}

// Navigation
function goToday() {{
  currentDay = calcToday();
  renderDay(currentDay);
}}
function prevDay() {{
  if (currentDay > 1) {{ currentDay--; renderDay(currentDay); }}
}}
function nextDay() {{
  if (currentDay < 365) {{ currentDay++; renderDay(currentDay); }}
}}
function jumpTo(n) {{
  n = parseInt(n);
  if (n >= 1 && n <= 365) {{ currentDay = n; renderDay(currentDay); }}
}}
function jumpToPhase(phase) {{
  const map = {{ Foundation: 1, Intermediate: 96, Advanced: 196, Integration: 291 }};
  jumpTo(map[phase] || 1);
}}

// Picker
function openPicker() {{
  const grid = document.getElementById("day-grid");
  grid.innerHTML = "";
  for (let d = 1; d <= 365; d++) {{
    const cell = document.createElement("div");
    cell.className = "day-cell";
    cell.textContent = d;
    const comp = dayCompletion(d);
    if (d === currentDay) cell.classList.add("current");
    else if (comp >= 1) cell.classList.add("completed");
    else if (comp > 0) cell.classList.add("partial");
    cell.onclick = () => {{ jumpTo(d); closePicker(); }};
    grid.appendChild(cell);
  }}
  document.getElementById("picker-overlay").classList.add("open");
}}
function closePicker() {{
  document.getElementById("picker-overlay").classList.remove("open");
}}

// Keyboard nav
document.addEventListener("keydown", e => {{
  if (e.target.tagName === "INPUT") return;
  if (e.key === "ArrowLeft") prevDay();
  if (e.key === "ArrowRight") nextDay();
  if (e.key === "t") goToday();
  if (e.key === "Escape") closePicker();
}});

// Init
currentDay = calcToday();
if (currentDay < 1 || currentDay > 365) currentDay = 1;
renderDay(currentDay);
updateStats();
</script>
</body>
</html>'''


def main():
    xlsx = "/home/ziad/git/claude_code_roadmap_for_fanngLevel_softwareEngineer/365_day_study_plan.xlsx"
    out = "/home/ziad/git/claude_code_roadmap_for_fanngLevel_softwareEngineer/study_dashboard.html"

    print("Extracting data from Excel...")
    data = extract_data(xlsx)
    print(f"Extracted {len(data)} days")

    print("Building HTML...")
    html = build_html(data)

    with open(out, "w", encoding="utf-8") as f:
        f.write(html)

    print(f"Done! Open in browser: {out}")
    print(f"File size: {len(html) // 1024} KB")


if __name__ == "__main__":
    main()
